# -*- coding: utf-8 -*-
"""Verifikasi struktur dokumen .docx yang dihasilkan."""
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PATH = r"C:\Users\HP\Documents\P UI 2\belajar-vue\Laporan_Praktikum_PUI_Lanjut_Modul0-1.docx"

doc = Document(PATH)
print("Paragraf          :", len(doc.paragraphs))
print("Tabel             :", len(doc.tables))

headings = [(p.style.name, p.text.strip()) for p in doc.paragraphs
            if p.style.name.startswith("Heading") and p.text.strip()]
print("Heading           :", len(headings))
for st, tx in headings:
    print("   ", st[-1], "|", tx)

with zipfile.ZipFile(PATH) as z:
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    print("Gambar tertanam   :", len(media))
    xml = z.read("word/document.xml").decode("utf-8")
    print("Field TOC         :", xml.count("TOC \\o"))
    print("Field PAGE        :", xml.count("PAGE"))
    print("updateFields      :", "updateFields" in z.read("word/settings.xml").decode("utf-8"))

# daftar isi (caption gambar & tabel)
caps = [p.text.strip() for p in doc.paragraphs
        if re.match(r"^(Gambar|Tabel|Kode)\s+\d", p.text.strip())]
print("\n--- Caption di dalam dokumen (%d) ---" % len(caps))
for c in caps:
    print("   ", c[:95])

# cek tabel: jumlah kolom & baris
print("\n--- Ringkasan tabel ---")
for i, t in enumerate(doc.tables, 1):
    hdr = " | ".join(c.text.strip()[:18] for c in t.rows[0].cells)
    print("  T%02d  %dr x %dc   %s" % (i, len(t.rows), len(t.columns), hdr))
