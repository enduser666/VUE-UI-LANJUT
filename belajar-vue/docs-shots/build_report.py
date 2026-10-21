# -*- coding: utf-8 -*-
"""Build Laporan Praktikum Pengembangan UI Lanjut (Modul 0 & 1)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from report_lib import Report, NAVY, PRIMARY, GREY, BLACK  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402
import sections_a  # noqa: E402
import sections_b  # noqa: E402

BASE = r"C:\Users\HP\Documents\P UI 2\belajar-vue"
SHOTS = os.path.join(BASE, "docs-shots")
OUT = os.path.join(BASE, "Laporan_Praktikum_PUI_Lanjut_Modul0-1.docx")

R = Report(OUT, images_dir=SHOTS)

# =============================================================== HALAMAN JUDUL
R.blank(2)
R.para("LAPORAN PRAKTIKUM", size=16, bold=True, align="center", indent=0,
       space_after=2, color=NAVY)
R.para("MATA KULIAH PENGEMBANGAN UI LANJUT", size=14, bold=True, align="center",
       indent=0, space_after=6, color=PRIMARY)
R.para("MODUL 0 & MODUL 1", size=13, bold=True, align="center", indent=0,
       space_after=16, color=NAVY)
R.para("INSTALASI, KONFIGURASI, SERTA STRUKTUR APLIKASI DAN NAVIGASI "
       "PADA VUE.JS", size=13, bold=True, align="center", indent=0,
       space_after=6, line=1.4)
R.para("(Studi Implementasi pada Aplikasi “Gatherly” — Platform Komunitas Event)",
       size=12, italic=True, align="center", indent=0, space_after=26, color=GREY)

t = R.doc.add_table(rows=0, cols=2)
t.alignment = 1
t.autofit = False
R._fixed_layout(t, [4.6, 8.4])
for k, v in [("Disusun oleh", ""), ("Nama", "Azriel Sabiq"),
             ("NIM", "102042300059"), ("Kelas", "SI-KH-23-05")]:
    cells = t.add_row().cells
    cells[0].width = Cm(4.6)
    cells[1].width = Cm(8.4)
    for cell, txt, bold, al in ((cells[0], k, False, 2), (cells[1], v, True, 0)):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = al
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_after = Pt(2)
        R._font(p.add_run(txt), "Times New Roman", 12, bold, False, BLACK)
R.blank(3)
R.para("PROGRAM STUDI S1 SISTEM INFORMASI", size=13, bold=True, align="center",
       indent=0, space_after=2, color=NAVY)
R.para("FAKULTAS REKAYASA INDUSTRI", size=13, bold=True, align="center", indent=0,
       space_after=2, color=NAVY)
R.para("UNIVERSITAS TELKOM", size=13, bold=True, align="center", indent=0,
       space_after=2, color=NAVY)
R.para("2026", size=13, bold=True, align="center", indent=0, space_after=0,
       color=NAVY)

# =============================================================== IDENTITAS
R.page_break()
R.para("DAFTAR IDENTITAS LAPORAN", size=13, bold=True, align="center", indent=0,
       space_after=14, color=NAVY)
R.table(
    ["Komponen", "Keterangan"],
    [["Mata Kuliah", "Pengembangan UI Lanjut"],
     ["Modul", "Modul 0 (Instalasi dan Konfigurasi Vue.js) dan Modul 1 "
               "(Struktur Aplikasi dan Navigasi dalam Vue.js)"],
     ["Referensi Utama", "Modul Praktikum — Modul 0 & 1 Pengembangan UI Lanjut, "
                         "Tim Asisten Praktikum, Tahun Ajaran 2026/2027"],
     ["Judul Proyek", "belajar-vue — Aplikasi Gatherly (Platform Komunitas Event)"],
     ["Lokasi Proyek", r"C:\Users\HP\Documents\P UI 2\belajar-vue"],
     ["Disusun oleh", "Azriel Sabiq — 102042300059 — SI-KH-23-05"],
     ["Tanggal Penyusunan", "2026"],
     ["Teknologi Utama", "Vue 3.5, Vue Router 5, Pinia 4, Vite 8, JavaScript ES Module"],
     ["Status Build", "Berhasil (npm run build selesai tanpa galat)"]],
    widths=[4.2, 10.8], size=10.5)

# =============================================================== DAFTAR ISI
R.page_break()
R.h1("DAFTAR ISI", page_break=False)
R.toc()

# =============================================================== DAFTAR GAMBAR
R.page_break()
R.h1("DAFTAR GAMBAR", page_break=False)
for no, cap in [
    ("Gambar 3.1", "Peta Situs (Site Map) Aplikasi Gatherly beserta Hierarki Parent–Child"),
    ("Gambar 3.2", "Arsitektur Lapisan Aplikasi Vue.js pada Proyek belajar-vue"),
    ("Gambar 3.3", "Perancangan Hierarki Navigasi pada Komponen Navbar.vue"),
    ("Gambar 3.4", "Pengelompokan Komponen Reusable pada Direktori src/components"),
    ("Gambar 3.5", "Alur Setup Environment Praktikum Modul 0"),
    ("Gambar 4.1", "Tampilan Dropdown dan Submenu pada Komponen Navbar.vue"),
    ("Gambar 4.2", "Tampilan Halaman Utama (Home.vue) pada Hero Section"),
    ("Gambar 4.3", "Tampilan Halaman About (About.vue)"),
    ("Gambar 4.4", "Tampilan Halaman Contact (Contact.vue) beserta formulir"),
    ("Gambar 4.5", "Tampilan Halaman Browse (Parent) dan Event List (Child)"),
    ("Gambar 4.6", "Tampilan Halaman Category (Category.vue)"),
    ("Gambar 4.7", "Tampilan Halaman Detail Event (EventDetail.vue)"),
    ("Gambar 4.8", "Tampilan Dashboard Organizer beserta metrik dan tabel rekap"),
    ("Gambar 4.9", "Tampilan Halaman Manajemen Event (EventManagement.vue)"),
    ("Gambar 4.10", "Tampilan Formulir Pembuatan Event Baru"),
    ("Gambar 4.11", "Tampilan Modal Konfirmasi Aksi pada Komponen AppModal.vue"),
]:
    R.para("%s   %s" % (no, cap), size=11, align="left", indent=0, space_after=4,
           line=1.3, left_indent=0.5)

# =============================================================== DAFTAR TABEL
R.page_break()
R.h1("DAFTAR TABEL", page_break=False)
for no, cap in [
    ("Tabel 2.1", "Perbandingan Konsep Single Page Application dan Multi Page Application"),
    ("Tabel 3.1", "Capaian Pembelajaran Mata Kuliah (CPMK) Modul 0 dan Modul 1"),
    ("Tabel 3.2", "Analisis Kebutuhan Fungsional Aplikasi Gatherly"),
    ("Tabel 3.3", "Kontrak Props dan Emits Antarkomponen"),
    ("Tabel 3.4", "Daftar Rute Aplikasi pada src/router/index.js"),
    ("Tabel 3.5", "Spesifikasi Teknis Proyek belajar-vue"),
    ("Tabel 4.1", "Konfigurasi Project pada Generator create-vue"),
    ("Tabel 4.2", "Verifikasi Environment Pengembangan pada Mesin Praktikan"),
    ("Tabel 4.3", "Perincian File Sumber pada Direktori src/"),
    ("Tabel 4.4", "Ringkasan Hasil Build Produksi (npm run build)"),
    ("Tabel 5.1", "Hasil Pengujian Black-Box Fungsional"),
    ("Tabel 5.2", "Pencapaian Capaian Pembelajaran Mata Kuliah"),
]:
    R.para("%s   %s" % (no, cap), size=11, align="left", indent=0, space_after=4,
           line=1.3, left_indent=0.5)

# =============================================================== BAB I
R.h1("BAB I  PENDAHULUAN")

R.h2("1.1  Latar Belakang")
R.para("Perkembangan teknologi antarmuka web telah mengubah cara manusia "
       "mengakses dan mengelola informasi. Aplikasi berbasis web modern tidak "
       "lagi sekadar menampilkan dokumen statis, melainkan menyediakan "
       "pengalaman interaktif yang dinamis, cepat, dan konsisten pada "
       "berbagai perangkat. Salah satu framework yang paling banyak digunakan "
       "dalam ekosistem tersebut adalah Vue.js.")
R.para("Vue.js merupakan framework front-end progresif yang dirancang oleh "
       "Evan You. Karakteristik utamanya terletak pada pendekatan reaktif dan "
       "deklaratif: pengembang cukup mendeskripsikan tampilan yang "
       "dikehendaki pada kondisi tertentu, sedangkan framework yang "
       "bertanggung jawab melacak perubahan state dan otomatis memperbarui "
       "antarmuka yang ditampilkan. Pada versi 3, Vue.js memperkenalkan "
       "Composition API berupa susunan setup() yang memberi keleluasaan "
       "pengorganisasian logika komponen sekaligus menutup keterbatasan "
       "Options API yang lazim digunakan pada Vue 2.")
R.para("Pada sisi toolchain, Vite berperan sebagai development server sekaligus "
       "build tool. Berbeda dengan bundler konvensional yang harus memproses "
       "seluruh berkas aplikasi sebelum halaman dapat disajikan, Vite "
       "memanfaatkan native ES Module sehingga proses start-up development "
       "server menjadi jauh lebih cepat. Kombinasi Vue.js dan Vite kemudian "
       "menjadi fondasi yang lazim dipakai pada praktikum pemrograman "
       "antarmuka web modern, termasuk pada mata kuliah Pengembangan UI "
       "Lanjut.")
R.para("Dalam konteks mata kuliah Pengembangan UI Lanjut, modul pertama "
       "(Modul 0) berfokus pada penyiapan environment pengembangan sehingga "
       "praktikan dapat membuat, menjalankan, dan menghentikan server lokal "
       "secara mandiri. Modul kedua (Modul 1) kemudian berfokus pada struktur "
       "aplikasi dan navigasi, yaitu penerjemahan konsep Site Map menjadi "
       "kode yang benar-benar berjalan melalui Vue Router, dengan penekanan "
       "pada hierarki Parent–Child, komponen Navbar sebagai alat "
       "transportasi pengguna, dan Breadcrumb sebagai penunjuk arah.")
R.para("Praktikum ini mendokumentasikan proses tersebut secara utuh, mulai "
       "dari instalasi Visual Studio Code, Node.js, dan npm, pembuatan "
       "kerangka project dengan generator create-vue, pengimplementasian "
       "hierarki Parent–Child pada src/router/index.js, hingga "
       "konstruksi komponen Navbar dan Breadcrumb beserta seluruh halaman "
       "destinasi pendukungnya. Sebagai kasus studi, project yang digunakan "
       "bernama belajar-vue dengan identitas produk Gatherly, sebuah platform "
       "komunitas event yang memadukan katalog publik dengan dasbor "
       "organizer.")

R.h2("1.2  Identifikasi Masalah")
R.para("Berdasarkan kondisi awal praktikan, beberapa persoalan yang umum "
       "ditemukan pada perkuliahan Pengembangan UI Lanjut adalah sebagai "
       "berikut.")
for x in [
    "Sebagian besar praktikan belum memiliki environment pengembangan yang "
    "terstandarisasi sehingga proses pembuatan project Vue.js tertunda karena "
    "kesalahan konfigurasi Node.js, npm, atau ekstensi editor.",
    "Pemahaman mengenai pemisahan peran antara View, Layout, dan Component "
    "masih lemah, sehingga penulisan kode cenderung menempatkan seluruh "
    "antarmuka pada satu berkas halaman.",
    "Belum adanya standar pemetaan Site Map ke tabel rute, yang mengakibatkan "
    "struktur navigasi berkembang tidak tertata dan sulit dipelihara.",
    "Tidak tersedianya mekanisme orientasi posisi pengguna pada halaman "
    "bernested, sehingga pengguna kehilangan konteks saat berpindah halaman.",
    "Belum diterapkannya reusable component, sehingga elemen antarmuka yang "
    "serupa ditulis berulang kali pada beberapa berkas.",
]:
    R.bullet(x)

R.h2("1.3  Rumusan Masalah")
R.para("Berdasarkan identifikasi masalah di atas, rumusan masalah pada "
       "praktikum ini adalah:")
for x in [
    "Bagaimana menyiapkan environment pengembangan yang lengkap agar "
    "project Vue.js dapat dibuat, dijalankan, dan dihentikan dengan benar?",
    "Bagaimana menerjemahkan konsep Site Map menjadi hierarki rute "
    "Parent–Child pada src/router/index.js?",
    "Bagaimana merancang komponen Navbar interaktif berbasis router-link "
    "beserta struktur dropdown dan submenu bertingkat?",
    "Bagaimana merancang komponen Breadcrumb sebagai Location Indicator "
    "yang otomatis mengikuti kedalaman rute yang sedang diakses?",
    "Bagaimana menyusun halaman-halaman destinasi (View) beserta komponen "
    "reusable agar struktur aplikasi tetap ringkas dan mudah dikembangkan?",
]:
    R.numbered(x)

R.h2("1.4  Tujuan")
R.para("Tujuan dari penyusunan laporan praktikum ini adalah:")
for x in [
    "Mendokumentasikan proses instalasi dan konfigurasi environment "
    "pengembangan Vue.js menggunakan Visual Studio Code, Node.js, npm, dan "
    "Vite.",
    "Membuat kerangka project Vue.js menggunakan generator create-vue dengan "
    "pilihan konfigurasi yang sesuai kebutuhan perkuliahan.",
    "Menerapkan hierarki Parent–Child pada konfigurasi Vue Router untuk "
    "menerjemahkan Site Map aplikasi Gatherly.",
    "Merancang dan mengimplementasikan komponen Navbar dan Breadcrumb "
    "sebagai mekanisme transisi dan orientasi pengguna.",
    "Membangun halaman-halaman destinasi beserta komponen reusable "
    "(AppButton, AppCard, AppModal, EventCard, SearchBar, CategoryFilter) "
    "dan mendokumentasikan hasil verifikasinya.",
]:
    R.bullet(x)

R.h2("1.5  Manfaat")
for x in [
    "Bagi praktikan: memberikan bekal pemahaman praktis mengenai instalasi, "
    "konfigurasi, dan struktur aplikasi Vue.js yang dapat langsung "
    "diterapkan pada modul-modul berikutnya.",
    "Bagi dosen pengampu: menjadi bahan evaluasi yang terukur karena "
    "setiap klaim pada laporan diverifikasi langsung pada aplikasi yang "
    "berjalan.",
    "Bagi pembaca: memperoleh gambaran utuh mengenai hubungan antara Site Map, "
    "View, Layout, dan Component.",
]:
    R.bullet(x)

R.h2("1.6  Batasan Masalah")
for x in [
    "Praktikum ini mencakup Modul 0 dan Modul 1 sesuai handout Modul 0 & 1 "
    "Pengembangan UI Lanjut; modul lanjutan seperti manajemen state "
    "berbasis Pinia dan pengujian otomatis tidak dibahas sebagai fokus "
    "utama, meskipun sebagian kapabilitasnya ikut digunakan.",
    "Pengujian dilakukan pada lingkungan lokal (localhost) menggunakan "
    "browser Chromium, tanpa pengujian lintas browser secara ekstensif.",
    "Data yang digunakan pada verifikasi adalah data contoh yang disimulasikan "
    "pada localStorage, bukan basis data server.",
]:
    R.bullet(x)

# =============================================================== BAB II - VI
sections_a.bab2(R)
sections_a.bab3(R)
sections_b.bab4(R)
sections_b.bab5(R)
sections_b.bab6(R)
sections_b.back_matter(R)

# =============================================================== SIMPAN
path = R.save()
print("Dokumen tersimpan:", path)
print("Ukuran: %d byte" % os.path.getsize(path))
