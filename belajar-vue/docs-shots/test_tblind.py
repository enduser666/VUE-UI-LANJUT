# -*- coding: utf-8 -*-
"""Uji isolasi: apakah insert_element_before untuk w:tblInd bekerja?"""
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

d = Document()
t = d.add_table(rows=1, cols=2)
tblPr = t._tbl.tblPr
print("sebelum:", [c.tag.split("}")[1] for c in tblPr])

el = OxmlElement("w:tblInd")
el.set(qn("w:w"), "0")
el.set(qn("w:type"), "dxa")
tblPr.insert_element_before(el, "w:tblBorders", "w:shd", "w:tblLayout",
                            "w:tblCellMar", "w:tblLook")
print("sesudah:", [c.tag.split("}")[1] for c in tblPr])

# uji dengan tabel yang sudah punya tblLayout
d2 = Document()
t2 = d2.add_table(rows=1, cols=2)
t2.autofit = False
tblPr2 = t2._tbl.tblPr
print("tbl2 sebelum:", [c.tag.split("}")[1] for c in tblPr2])
el2 = OxmlElement("w:tblInd")
el2.set(qn("w:w"), "0")
el2.set(qn("w:type"), "dxa")
tblPr2.insert_element_before(el2, "w:tblBorders", "w:shd", "w:tblLayout",
                             "w:tblCellMar", "w:tblLook")
print("tbl2 sesudah:", [c.tag.split("}")[1] for c in tblPr2])
